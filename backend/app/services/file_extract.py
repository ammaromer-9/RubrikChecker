"""Extract plain text from uploaded PDF, .docx and .txt files."""

from __future__ import annotations

import io
from pathlib import PurePath

import docx
import pdfplumber


class FileExtractionError(Exception):
    """Base class; `status_code` is the HTTP status the API should return."""

    status_code = 422


class FileTooLargeError(FileExtractionError):
    status_code = 413


class UnsupportedFileTypeError(FileExtractionError):
    status_code = 415


class NoTextFoundError(FileExtractionError):
    """The file parsed but contained no text (e.g. a scanned PDF)."""


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text(filename: str, content: bytes, *, max_bytes: int) -> str:
    if len(content) > max_bytes:
        raise FileTooLargeError(f"File exceeds the {max_bytes // (1024 * 1024)} MB upload limit.")

    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        text = _extract_pdf(content)
    elif extension == ".docx":
        text = _extract_docx(content)
    elif extension == ".txt":
        text = _decode_text(content)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type {extension or '(none)'}. Upload a PDF, .docx or .txt file."
        )

    text = text.strip()
    if not text:
        raise NoTextFoundError(
            "No text could be extracted. If this is a scanned PDF, paste the text instead."
        )
    return text


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception as exc:
        raise FileExtractionError("Could not read this PDF. It may be corrupted.") from exc
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise FileExtractionError("Could not read this .docx file. It may be corrupted.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    # Rubrics are often tables, so include cell text too.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)
